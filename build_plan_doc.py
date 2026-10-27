from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'output' / '海南省电机监控系统一期建设方案.docx'
IMG = ROOT / 'output' / 'ui-regional'
MAIN = ROOT / 'design' / 'monitor-initial.png'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6); styles['Normal'].paragraph_format.line_spacing = 1.15
for name, size, color in [('Title', 24, '000000'), ('Heading 1', 17, '000000'), ('Heading 2', 13, '000000'), ('Heading 3', 11, '000000')]:
    st = styles[name]; st.font.name = 'Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(13 if name != 'Title' else 0); st.paragraph_format.space_after = Pt(6)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color='D9E1EA'):
    tblPr = table._tbl.tblPr; b = tblPr.first_child_found_in('w:tblBorders')
    if b is None: b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge; el = b.find(qn(tag))
        if el is None: el = OxmlElement(tag); b.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),color)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'17365D'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx%2==1: shade(cells[i],'F3F6FA')
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def p(text='', bold_prefix=None):
    para=doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold=True; para.add_run(text[len(bold_prefix):])
    else: para.add_run(text)
    return para

def bullet(text, level=0):
    para=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); para.add_run(text); return para

def image(path, caption, width=6.4):
    if not Path(path).exists(): return
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.add_run().add_picture(str(path), width=Inches(width))
    cp=doc.add_paragraph(caption); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in cp.runs: r.font.size=Pt(9); r.font.italic=True; r.font.color.rgb=RGBColor(89,89,89)

# Title page
doc.add_paragraph('海南省电机监控系统一期建设方案', style='Title')
sub=doc.add_paragraph('面向约 1000 台电机的云端采集与本地大屏监控'); sub.runs[0].font.size=Pt(14); sub.runs[0].font.color.rgb=RGBColor(31,78,121)
doc.add_paragraph('版本：V1.0 方案讨论稿')
doc.add_paragraph('编制日期：2026 年 9 月 19 日')
doc.add_paragraph('适用对象：项目负责人、甲方业务代表、厂商接口对接人员')
doc.add_paragraph()
p('本方案用于在甲方提供厂商设备接口文档、测试账号和设备资料前，统一一期建设范围、界面方向、技术边界和验收方法。厂商未确认的字段、调用频率、故障码、历史能力和授权条件，均保留为待核实事项，不作为采购、工期或性能承诺。')
doc.add_page_break()

doc.add_heading('一 项目结论与待确认事项', 1)
p('一期建议建设“本地服务器采集与存储 + 局域网上位机展示”的单机系统。系统从厂商官方云端读取设备资料和运行状态，在本地保存历史、处理报警，并向海南省电机运行监控大屏提供统一数据。浏览器不直接接触厂商密钥，也不提供远程启停、复位或参数修改。')
table(['类别','当前方案','状态'],[
    ['设备规模','初期约 1000 台电机','已作为方案基线'],['数据来源','厂商官方云端接口','等待甲方提供资料'],['展示方式','海南省地图总览 + 区域设备分屏 + 单机详情','UI 方向已形成'],['报警方式','故障事件与通信状态分开，按需提示，支持确认和恢复','规则等待厂商确认'],['历史数据','原始遥测、聚合趋势、故障事件分层保存','保存期限待审批'],['控制边界','只监控和报警，不下发设备控制指令','一期固定边界'],['部署位置','一台本地服务器、一台上位机和大屏','硬件环境待现场确认']])
doc.add_heading('甲方需补充的资料', 2)
for x in ['厂商接口文档、测试账号、设备清单和故障码解释。','厂商设备唯一标识、所属人、乡镇、村庄、位置及基本参数字段说明。','电流、电压、温度、转速的测点含义、相别、单位和时间戳规则。','接口调用限制、批量方式、令牌续期、历史查询、补采能力和本地保存授权。','大屏类型、尺寸、有效分辨率、宽高比、观看距离、输入方式、运行时长和现有上位机条件。','历史保存期限、备份恢复目标、用户角色和报警确认责任人。']:
    bullet(x)

doc.add_heading('二 建设目标与范围', 1)
doc.add_heading('建设目标', 2)
for x in ['在一屏内掌握海南省各市县电机运行概况、故障分布和数据新鲜度。','按市县、乡镇和村庄逐级进入区域设备页，保持设备编号稳定排序。','点击单台设备查看当前参数、最后有效数据、历史趋势、故障和通信事件。','故障确认不清除故障，设备恢复单独记录；数据中断不误判为电机故障。','服务器重启、浏览器重连和公网中断后能够恢复或明确标识数据状态。']:
    bullet(x)
doc.add_heading('一期包含', 2)
table(['模块','内容'],[
    ['省级总览','海南地图、市县聚合运行数量、接入/运行/故障指标、24 小时运行趋势、历史健康状态、按需报警提示'],
    ['区域设备','区域树、区域统计、关键词与状态组合筛选、分页设备卡片、轮播暂停和恢复'],
    ['设备详情','设备资料、当前/最后有效读数、历史趋势、数据质量、故障与通信事件'],
    ['报警中心','活动报警、未确认数量、确认记录、恢复记录、服务重启恢复'],
    ['后台服务','云端适配、采集调度、数据校验、持久化、聚合、Web API、运行日志和备份']])
doc.add_heading('一期不包含', 2)
p('地图编辑、人员档案维护、短信或手机推送、声音报警、远程启停、远程复位、参数修改、独立地图 GIS 管理、未经审批的机器学习诊断和多租户平台化建设。')

doc.add_heading('三 界面方案', 1)
p('界面采用深蓝底、青色边界和克制的珊瑚色异常标记。主屏是远距离查看的总览，不铺开单台设备明细；区域设备和单机详情作为操作分屏。报警没有故障时不占据固定大窗口，有新事件时以提示条和数量出现，已确认但未恢复的故障仍保留在地图标记和统计中。')
image(MAIN, '图 1 海南省电机运行监控中心总览设计稿，地图为示意数据。')
doc.add_heading('主屏信息层级', 2)
table(['层级','展示内容','设计原则'],[
    ['第一层','海南地图、市县运行数量、故障位置、三个关键指标','先回答哪里需要关注'],['第二层','近 24 小时运行趋势、近 7 天故障/离线历史趋势','回答状态是否持续变化'],['第三层','按需出现的报警提示条和导航数量','有事件才突出'],['第四层','区域设备与单机详情入口','主动操作后展开']])
doc.add_heading('区域设备页面', 2)
image(IMG/'01-region-overview.png', '图 2 区域设备总览：区域树、六台简要卡片和组合筛选。')
p('区域设备页每页默认展示 6 台，卡片只显示名称、编号、状态、所属人和位置。电流、电压、温度、转速和故障代码进入单机详情。卡片保持设备编号顺序，不因报警重排。区域统计和筛选结果分别显示，避免“搜索无结果”被误解为区域没有设备。')
image(IMG/'03-search-empty.png', '图 3 搜索无结果状态：区域总数保持，筛选结果明确为 0。')
doc.add_heading('单机详情页面', 2)
image(IMG/'02-device-detail.png', '图 4 故障设备详情：故障已确认但尚未恢复，展示 24 小时温度趋势。')
image(IMG/'04-offline-detail.png', '图 5 离线设备详情：最后有效读数与无新数据区间明确区分。')
doc.add_heading('界面状态规则', 2)
table(['状态','页面表现','禁止的推断'],[
    ['运行','状态为运行，显示最新有效读数','不能仅凭绿色推断全部数据完整'],['停止','显示已停止，保留通信和数据新鲜度','停止不等于故障'],['故障','地图和卡片标红，记录故障码和时间','确认不等于恢复'],['离线','显示最后有效数据时间，实时读数变为未知','离线不等于停止或故障'],['数据过期','显示延迟/过期提示，保留原状态','不能把过期统一判为离线']])

doc.add_heading('四 总体技术架构', 1)
p('推荐数据流为：电机与采集设备 → 厂商云平台 → 本地服务器云端适配器 → 采集调度与校验 → 数据库和报警处理 → 本地 Web API/实时通知 → 局域网上位机与大屏。厂商差异集中在适配模块，上层使用统一字段。')
table(['组件','职责','一期建议'],[
    ['云端适配器','鉴权、分页、批量拉取、字段转换、限流退避','按厂商资料单独实现'],['采集调度','按周期分批请求、记录批次、失败重试','初步建议 20 秒，接口验证后确定'],['数据层','当前状态、原始遥测、聚合数据、报警、日志','单机部署，数据库选型待容量试采'],['报警处理','故障实例去重、确认、恢复、通信级报警','按故障码和事件 ID 实现'],['Web 服务','区域查询、设备分页、详情、历史和报警接口','只向上位机开放'],['展示端','主屏、区域设备、详情、报警记录','浏览器访问本地服务']])
doc.add_heading('是否采用 Kafka', 2)
p('当前 1000 台设备和单台本地服务器不要求一期直接引入 Kafka 集群。先用单机采集服务、持久化数据库和进程内或轻量队列完成 20–50 台试点；当厂商数据量、并发客户端、补采规模或可用性要求证明需要时，再评估 Kafka。开源工业 IoT 项目可参考其服务分层、WebSocket 推送、报警引擎和监控指标，但其模拟 Producer 不能替代厂商云适配器。')

doc.add_heading('五 数据模型与接口计划', 1)
doc.add_heading('统一数据对象', 2)
table(['对象','核心字段','说明'],[
    ['设备资料','内部 ID、厂商设备 ID、名称、所属人、乡镇、村庄、位置、基本参数、资料同步时间','当前资料与报警发生时归属摘要分开保存'],['当前状态','运行状态、通信状态、数据新鲜度、测点值与单位、故障码、测量/接收时间','运行、故障、离线和过期独立表达'],['历史遥测','设备 ID、测量时间、接收时间、值、单位、质量标记、来源序列号','重复数据按稳定键去重'],['报警记录','报警 ID、设备/来源、故障码、发生、确认、恢复、状态','持续故障只保留一个活动实例'],['运行日志','采集批次、成功数、失败摘要、最近成功时间','区分接口整体异常与设备故障']])
doc.add_heading('首批本地 API 草案', 2)
table(['接口','用途','厂商资料依赖'],[
    ['GET /api/regions/tree','区域树和数量','区域字段'],['GET /api/regions/{id}/summary','区域统计和更新时间','设备列表/状态字段'],['GET /api/devices','分页、搜索、状态筛选','设备 ID、名称、归属、状态'],['GET /api/devices/{id}','资料、当前读数、故障和通信状态','全部设备字段'],['GET /api/devices/{id}/telemetry','历史趋势','测量时间、测点、历史期限'],['GET /api/alarms','活动/未确认报警','故障事件和恢复事件'],['POST /api/alarms/{id}/ack','记录确认人和时间','本地权限与审计']])
p('接口命名和字段类型为本地服务草案，等待厂商文档后形成字段映射表、错误码表和契约测试。')

doc.add_heading('六 报警与异常处理', 1)
p('正式故障以厂商明确故障码或故障事件为依据，不把演示中的温度阈值直接当作生产规则。设备故障、通信中断、云接口整体不可用、部分返回和数据过期必须分别处理。')
table(['场景','处理方式'],[
    ['新故障','建立活动报警，地图按市县聚合，区域和设备页保留标记'],['持续故障','同一设备、同一故障码只保留一个活动实例'],['人工确认','写入确认人和时间；关闭提示不等于确认'],['故障恢复','依据厂商恢复事件或审批后的连续正常采样规则归档'],['设备离线','显示最后有效数据时间，不将最后值当作当前值'],['云接口不可用','产生来源级通信报警，不批量制造电机故障'],['部分返回','标记采集批次不完整，暂不批量判离线'],['服务重启','恢复活动报警、当前状态和未确认摘要']])

doc.add_heading('七 存储与部署建议', 1)
p('建议先用真实样本试采 7 天，再决定磁盘、聚合粒度和备份策略。以下为讨论基线，不是采购承诺。')
table(['数据','初步建议','需确认'],[
    ['原始遥测','90 天','采样频率、样本大小、是否保存原文'],['1 分钟聚合','1 年','历史查询需求'],['15 分钟聚合','3 年','是否纳入一期'],['故障与确认','5 年','合规与业务要求'],['当前设备资料','在用期间','退役和个人信息清理'],['备份','每日备份、保留近 30 天恢复点','可接受的数据损失窗口']])
table(['设备','起步建议','现场核查'],[
    ['服务器','8 核 CPU、32 GB 内存、约 2 TB 可用 SSD、UPS','真实样本容量、磁盘镜像、备份介质'],['上位机','按大屏原生分辨率稳定输出','显卡、接口、线缆、拼接/LED 控制器'],['网络','服务器稳定访问公网，上位机与服务器同局域网','固定地址、防火墙、时间同步'],['账户','查看、值班确认、系统管理三类角色','确认权限和审计要求']])

doc.add_heading('八 实施计划', 1)
table(['阶段','主要工作','输出与通过条件'],[
    ['0 方案确认','确认 UI 主屏、区域设备、详情边界；确定甲方接口对接人','方案意见记录，未决项有责任方'],['1 接口核验','取得接口文档、账号、设备清单；核对字段、单位、频率、限流、故障事件','接口能力核对表和样例字段映射'],['2 数据服务','实现适配、采集、校验、当前状态、日志和持久化','测试数据可采集，重复/乱序/失败可追踪'],['3 UI 对接','实现总览、区域设备、详情、历史和按需报警提示','固定 JSON 与真实接口均可切换，核心交互通过'],['4 小范围试点','20–50 台跨市县设备，真实大屏连续试运行 7 天','数据比对、断网、报警、重启、容量和观看距离评审'],['5 全量部署','扩展至约 1000 台，配置备份、监控、账户和培训','全量验收、部署交接和操作培训']])

doc.add_heading('九 验收建议', 1)
for x in ['设备唯一标识不重复，云端归属和位置字段正确，缺失字段按约定显示。','主屏在实际大屏原生分辨率和观看距离下可读，地图、报警和关键数字无遮挡。','区域、关键词、状态组合筛选正确；分页无漏项、重复或错误总数。','报警确认、恢复、重复故障和服务重启均可追溯；数据中断不误报电机故障。','厂商条件满足时，测量到展示延迟 P95 目标不超过 30 秒；在线报警到上位机提示 P95 目标不超过 5 秒。','验证限流、鉴权失效、部分返回、重复乱序、公网中断、浏览器重连和上位机重启。','连续 24 小时规模运行，核对有效唯一样本无静默丢失；保存和恢复达到批准目标。','客户端不暴露厂商密钥，系统无设备控制按钮和下发接口，确认操作有权限和审计。']:
    bullet(x)

doc.add_heading('十 当前决策与下一步', 1)
p('当前可以冻结 UI 结构和本地接口边界，不能冻结厂商字段、采集周期、历史期限、硬件规格和最终性能指标。待甲方提供厂商接口文档后，优先完成设备列表字段映射和分页方式，再依次接入设备详情、历史数据、报警事件和恢复事件。')
table(['优先级','待办','责任建议'],[
    ['P0','甲方提供接口文档、测试账号、设备清单、故障码表','甲方/厂商'],['P0','确定接口对接人和授权范围','甲方'],['P1','完成字段映射和接口能力核对','项目组/厂商'],['P1','用固定 JSON 先完成区域设备交互验收','项目组/甲方'],['P2','确定历史保存、备份恢复和硬件预算','甲方/项目组'],['P2','安排 20–50 台试点与实际大屏评审','项目组/现场']])
p('本方案可作为接口资料到位前的建设和评审基础。接口资料到位后，应形成 V1.1 字段映射、采集策略、部署清单和验收测试用例。')

# Footer
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('海南省电机监控系统一期建设方案  |  方案讨论稿'); run.font.size=Pt(8); run.font.color.rgb=RGBColor(120,120,120)

doc.core_properties.title='海南省电机监控系统一期建设方案'
doc.core_properties.subject='生产方案、界面设计与厂商接口待确认事项'
doc.core_properties.author='项目组'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
